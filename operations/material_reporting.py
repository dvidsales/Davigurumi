from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO, StringIO
import csv
from django import forms
from materials.models import Material
from portability.tables import safe_cell
from portability.views import download
from .reporting import PeriodForm


class MaterialPeriodForm(PeriodForm):
    material = forms.ModelChoiceField(
        label="Material",
        queryset=Material.objects.none(),
        required=False,
        empty_label="Todos, inclusive arquivados",
    )

    def __init__(self, *args, owner, **kwargs):
        super().__init__(*args, owner=owner, **kwargs)
        for name in ("client", "project", "status"):
            self.fields.pop(name)
        self.fields["material"].queryset = Material.objects.filter(owner=owner)


def movement_rows(movements):
    for entry in movements:
        yield [
            str(entry.pk),
            safe_cell(entry.material.name),
            entry.material.unit,
            entry.created_at.isoformat(),
            entry.get_kind_display(),
            entry.quantity,
            entry.unit_cost,
            (
                (entry.quantity * entry.unit_cost).quantize(
                    Decimal(".01"), rounding=ROUND_HALF_UP
                )
                if entry.unit_cost is not None
                else None
            ),
            "Conhecido" if entry.unit_cost is not None else "Desconhecido",
            safe_cell(entry.reason),
            str(entry.reverses_id) if entry.reverses_id else "",
        ]


def export_movements(movements, extension):
    headers = [
        "movimento",
        "material",
        "unidade",
        "data",
        "tipo",
        "quantidade_assinada",
        "custo_unitario_historico",
        "variacao_valor_estoque",
        "situacao_custo",
        "motivo",
        "compensa_movimento",
    ]
    if extension == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "Movimentos"
        sheet.append(headers)
        for row in movement_rows(movements):
            sheet.append(row)
            sheet.cell(sheet.max_row, 6).number_format = "0.000000"
            sheet.cell(sheet.max_row, 7).number_format = "0.000000"
            sheet.cell(sheet.max_row, 8).number_format = "0.00"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        output = BytesIO()
        book.save(output)
        return download(
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "movimentos-materiais.xlsx",
        )
    output = StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(headers)
    writer.writerows(movement_rows(movements))
    return download(
        "\ufeff" + output.getvalue(),
        "text/csv; charset=utf-8",
        "movimentos-materiais.csv",
    )
