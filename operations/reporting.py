from collections import defaultdict
from decimal import Decimal
from django import forms
from django.db.models import Sum
from django.utils import timezone
from finance.models import Payment, Refund
from materials.models import Material, StockReservation
from production.models import Order, ProductionSession
from sales.models import QuoteVersion, Client
from projects.models import Project


class PeriodForm(forms.Form):
    status = forms.ChoiceField(
        label="Produção",
        required=False,
        choices=[("", "Todas")] + Order._meta.get_field("status").choices,
    )
    client = forms.ModelChoiceField(
        label="Cliente",
        queryset=Client.objects.none(),
        required=False,
        empty_label="Todos",
    )
    project = forms.ModelChoiceField(
        label="Projeto",
        queryset=Project.objects.none(),
        required=False,
        empty_label="Todos",
    )

    def __init__(self, *args, owner, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["client"].queryset = Client.objects.filter(owner=owner)
        self.fields["project"].queryset = Project.objects.filter(owner=owner)

    start = forms.DateField(
        label="De", widget=forms.DateInput(format="%Y-%m-%d", attrs={"type": "date"})
    )
    end = forms.DateField(
        label="Até", widget=forms.DateInput(format="%Y-%m-%d", attrs={"type": "date"})
    )

    def clean(self):
        data = super().clean()
        if data.get("start") and data.get("end"):
            if data["end"] < data["start"]:
                raise forms.ValidationError(
                    "A data final deve ser igual ou posterior à inicial."
                )
            if (data["end"] - data["start"]).days > 3660:
                raise forms.ValidationError("Escolha um período de até 10 anos.")
        return data


def cash_period(owner, start, end):
    received = Payment.objects.filter(owner=owner, date__range=(start, end)).aggregate(
        total=Sum("amount")
    )["total"] or Decimal(0)
    refunded = Refund.objects.filter(
        allocation__payment__owner=owner, date__range=(start, end)
    ).aggregate(total=Sum("amount"))["total"] or Decimal(0)
    return {"receipts": received, "refunds": refunded, "net": received - refunded}


def replenishment(owner):
    """Remaining requirements subtract actual consumption and this item's reserves.
    Compare the sum of uncovered needs with free stock (already excludes every reserve).
    """
    from materials.stock import net_consumed

    required = defaultdict(Decimal)
    orders = Order.objects.filter(
        owner=owner, status__in=["waiting", "in_progress", "paused"]
    ).prefetch_related("items__consumptions__movement__compensations")
    reserved_totals = {
        (str(row["reference"]), str(row["layer__material_id"])): row["total"]
        for row in StockReservation.objects.filter(
            layer__material__owner=owner, remaining__gt=0
        )
        .values("reference", "layer__material_id")
        .annotate(total=Sum("remaining"))
    }
    for order in orders:
        for item in order.items.all():
            if item.retired:
                continue
            requirements = defaultdict(Decimal)
            for line in item.snapshot.get("materials", []):
                requirements[line["material"]] += Decimal(line["quantity"])
            for material_id, quantity in requirements.items():
                consumed = sum(
                    (
                        net_consumed(entry.movement)
                        for entry in item.consumptions.all()
                        if str(entry.movement.material_id) == material_id
                    ),
                    Decimal(0),
                )
                reserved = reserved_totals.get((str(item.pk), material_id), Decimal(0))
                required[material_id] += max(Decimal(0), quantity - consumed - reserved)
    rows = []
    for material in Material.objects.filter(
        owner=owner, pk__in=required
    ).prefetch_related("movements", "layers"):
        shortage = max(
            Decimal(0), required[str(material.pk)] - material.available_stock
        )
        if shortage:
            rows.append(
                {
                    "material": material,
                    "quantity": shortage,
                    "required": required[str(material.pk)],
                }
            )
    return rows


def overview(owner):
    today = timezone.localdate()
    orders = (
        Order.objects.filter(owner=owner)
        .select_related("current_version")
        .prefetch_related("allocations__refunds")
    )
    active = orders.filter(status__in=["waiting", "in_progress", "paused"])
    balances = sum(
        (order.balance for order in orders.exclude(status="cancelled")), Decimal(0)
    )
    return {
        "active_orders": active.count(),
        "late_orders": active.filter(production_due__lt=today).count(),
        "open_balance": balances,
        "pending_quotes": QuoteVersion.objects.filter(
            quote__owner=owner, status="sent", expires_at__gt=timezone.now()
        ).count(),
        "active_session": ProductionSession.objects.filter(
            owner=owner, ended_at__isnull=True
        )
        .select_related("item__order")
        .first(),
        "cash": cash_period(owner, today.replace(day=1), today),
        "shortages": replenishment(owner),
        "below_minimum": below_minimum(owner),
    }


REPORT_HEADERS = [
    "pedido",
    "cliente",
    "criacao",
    "producao",
    "entrega",
    "total_atual",
    "recebido_liquido_historico",
    "saldo_atual",
    "credito_atual",
]


def order_report_rows(orders):
    from portability.tables import safe_cell

    for order in orders:
        client = order.approved_version.quote.client
        yield [
            str(order.pk),
            safe_cell(client.name) if client else "",
            order.created_at.isoformat(),
            order.get_status_display(),
            order.get_delivery_status_display(),
            order.total,
            order.net_received,
            order.balance,
            order.credit,
        ]


def export_order_report(orders, format):
    from portability.views import download
    from io import BytesIO, StringIO

    if format == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "Pedidos"
        sheet.append(REPORT_HEADERS)
        for row in order_report_rows(orders):
            sheet.append(row)
            for cell in sheet[sheet.max_row][5:]:
                cell.number_format = "#,##0.00"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in {
            "A": 38,
            "B": 28,
            "C": 34,
            "D": 18,
            "E": 22,
            "F": 18,
            "G": 30,
            "H": 18,
            "I": 18,
        }.items():
            sheet.column_dimensions[column].width = width
        output = BytesIO()
        book.save(output)
        return download(
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "pedidos.xlsx",
        )
    import csv

    output = StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(REPORT_HEADERS)
    writer.writerows(order_report_rows(orders))
    return download(
        "\ufeff" + output.getvalue(), "text/csv; charset=utf-8", "pedidos.csv"
    )


def below_minimum(owner):
    rows = []
    for material in Material.objects.filter(
        owner=owner, minimum_stock__gt=0, is_archived=False
    ).prefetch_related("movements", "layers"):
        available = material.available_stock
        if available < material.minimum_stock:
            rows.append(
                {
                    "material": material,
                    "available": available,
                    "quantity": material.minimum_stock - available,
                }
            )
    return rows
