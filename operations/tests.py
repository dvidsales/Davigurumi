import uuid
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.core import mail
from .models import Notification, NotificationPreference, OutboxEvent
from .services import emit, deliver_outbox


class NotificationTests(TestCase):
    @override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
    def test_external_failure_keeps_event_and_retries_without_duplicate_internal_notification(
        self,
    ):
        owner = get_user_model().objects.create_user(
            username="notify", email="notify@example.test"
        )
        NotificationPreference.objects.create(owner=owner, email_enabled=True)
        target = uuid.uuid4()
        emit(
            owner=owner,
            kind="approved",
            object_id=target,
            message="Orçamento aprovado.",
        )
        emit(
            owner=owner,
            kind="approved",
            object_id=target,
            message="Orçamento aprovado.",
        )
        self.assertEqual(Notification.objects.count(), 1)
        self.assertEqual(OutboxEvent.objects.count(), 1)
        with patch(
            "operations.services.send_mail", side_effect=OSError("synthetic failure")
        ):
            self.assertEqual(deliver_outbox(), {"sent": 0, "failed": 1})
        event = OutboxEvent.objects.get()
        self.assertEqual(event.status, "pending")
        self.assertEqual(event.attempts, 1)
        from django.utils import timezone

        event.next_attempt = timezone.now()
        event.save()
        self.assertEqual(deliver_outbox(), {"sent": 1, "failed": 0})
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(deliver_outbox(), {"sent": 0, "failed": 0})


class ReportExportTests(TestCase):
    def setUp(self):
        from sales.models import Client, Quote, QuoteVersion, QuoteItem
        from projects.models import Project, ProjectRevision
        from production.models import Order

        self.owner = get_user_model().objects.create_user(
            username="report", email="report@example.test"
        )
        self.other = get_user_model().objects.create_user(
            username="report_other", email="report_other@example.test"
        )
        self.buyer = Client.objects.create(
            owner=self.owner, name='=HYPERLINK("https://example.test")'
        )
        self.foreign = Client.objects.create(owner=self.other, name="Outro cliente")
        self.project = Project.objects.create(owner=self.owner, name="Coelho")
        revision = ProjectRevision.objects.create(
            project=self.project, number=1, name="Coelho"
        )
        quote = Quote.objects.create(owner=self.owner, number=1, client=self.buyer)
        version = QuoteVersion.objects.create(quote=quote, number=1, total=25)
        for i in range(2):
            QuoteItem.objects.create(
                version=version,
                project_revision=revision,
                quantity=1,
                description="Peça",
                snapshot={},
                calculated_price=10,
                total=10,
            )
        self.order = Order.objects.create(
            owner=self.owner, approved_version=version, current_version=version
        )
        another = Quote.objects.create(owner=self.owner, number=2)
        another_version = QuoteVersion.objects.create(quote=another, number=1, total=30)
        Order.objects.create(
            owner=self.owner,
            approved_version=another_version,
            current_version=another_version,
            status="cancelled",
        )
        self.client.force_login(self.owner)

    def test_xlsx_filters_distinct_orders_and_escapes_formula_text(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from django.urls import reverse

        response = self.client.get(
            reverse("operations:reports"),
            {
                "export": "xlsx",
                "client": self.buyer.pk,
                "project": self.project.pk,
                "status": "waiting",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 1).value, str(self.order.pk))
        self.assertTrue(sheet.cell(2, 2).value.startswith("'="))
        self.assertEqual(sheet.cell(2, 2).data_type, "s")
        self.assertEqual(sheet.cell(2, 6).value, 25)
        self.assertEqual(sheet.cell(2, 8).value, 25)
        self.assertEqual(sheet.freeze_panes, "A2")

    def test_foreign_client_cannot_be_selected_and_invalid_dates_do_not_export(self):
        from django.urls import reverse

        url = reverse("operations:reports")
        response = self.client.get(url, {"export": "xlsx", "client": self.foreign.pk})
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/html", response["Content-Type"])
        self.assertNotContains(response, self.foreign.name)
        response = self.client.get(
            url, {"export": "xlsx", "start": "2026-10-07", "end": "2026-10-01"}
        )
        self.assertContains(response, "A data final")
        self.assertIn("text/html", response["Content-Type"])


class MinimumStockTests(TestCase):
    def test_reservations_reduce_available_threshold_and_daily_alerts_are_unique(self):
        from decimal import Decimal as D
        from materials.models import Material
        from materials.stock import receive_stock, reserve_stock
        from .reporting import below_minimum
        from .services import generate_alerts
        from datetime import date

        owner = get_user_model().objects.create_user(
            username="minimum", email="minimum@example.test"
        )
        other = get_user_model().objects.create_user(
            username="minimum_other", email="minimum_other@example.test"
        )
        material = Material.objects.create(
            owner=owner, name="Fio", kind="yarn", unit="g", minimum_stock=D("80")
        )
        Material.objects.create(
            owner=other, name="Privado", kind="yarn", unit="g", minimum_stock=D("100")
        )
        receive_stock(
            owner=owner,
            material_id=material.pk,
            quantity=D("100"),
            unit_cost=D(".1"),
            key=uuid.uuid4(),
            reason="Teste",
        )
        self.assertEqual(below_minimum(owner), [])
        reserve_stock(
            owner=owner,
            material_id=material.pk,
            quantity=D("30"),
            key=uuid.uuid4(),
            label="Trabalho",
        )
        rows = below_minimum(owner)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["available"], D("70"))
        self.assertEqual(rows[0]["quantity"], D("10"))
        self.assertEqual(generate_alerts(owner, today=date(2026, 10, 7)), 1)
        self.assertEqual(generate_alerts(owner, today=date(2026, 10, 7)), 0)
        self.assertEqual(generate_alerts(owner, today=date(2026, 10, 8)), 1)
        self.assertEqual(material.physical_stock, D("100"))
        material.minimum_stock = 0
        material.save(update_fields=["minimum_stock"])
        self.assertEqual(below_minimum(owner), [])
        self.assertEqual(generate_alerts(owner, today=date(2026, 10, 9)), 0)

    def test_material_form_keeps_optional_defaults_and_rejects_invalid_piece_minimum(
        self,
    ):
        from materials.forms import MaterialForm

        payload = {
            "name": "Botão",
            "kind": "accessory",
            "unit": "un",
            "initial_quantity": "0",
            "request_key": uuid.uuid4(),
        }
        form = MaterialForm(payload)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["minimum_stock"], 0)
        for value in ["-1", "1.5", "1000000"]:
            with self.subTest(value=value):
                self.assertFalse(
                    MaterialForm({**payload, "minimum_stock": value}).is_valid()
                )
        form = MaterialForm(
            {
                **payload,
                "minimum_stock": "10",
                "composition": "Algodão",
                "thickness": "Médio",
                "recommended_hook": "2,5 mm",
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        material = form.save(commit=False)
        self.assertEqual(material.composition, "Algodão")
        self.assertEqual(material.minimum_stock, 10)
