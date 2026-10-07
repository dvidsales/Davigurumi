import csv, json
from io import BytesIO, StringIO
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST
from .forms import TableImportForm, ArchiveForm
from .models import ImportJob
from .tables import preview_import, confirm_import, export_material_rows
from .archive import export_archive, import_archive


def download(body, content_type, name):
    response = HttpResponse(body, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{name}"'
    response["Cache-Control"] = "private, no-store"
    return response


@never_cache
@login_required
def index(request):
    form = TableImportForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            job = preview_import(
                owner=request.user,
                upload=form.cleaned_data["upload"],
                mapping=form.mapping(),
                locale=form.cleaned_data["locale"],
                separator=form.cleaned_data["separator"],
                sheet=form.cleaned_data["sheet"],
            )
        except ValidationError as exc:
            form.add_error(None, exc)
        else:
            return redirect("portability:preview", job_id=job.pk)
    return render(
        request,
        "portability/index.html",
        {
            "form": form,
            "jobs": ImportJob.objects.filter(owner=request.user).order_by(
                "-created_at"
            )[:20],
        },
    )


@never_cache
@login_required
def preview(request, job_id):
    job = get_object_or_404(ImportJob, pk=job_id, owner=request.user)
    return render(request, "portability/preview.html", {"job": job})


@require_POST
@login_required
def confirm(request, job_id):
    try:
        count = confirm_import(owner=request.user, job_id=job_id)
    except ValidationError as exc:
        messages.error(request, " ".join(exc.messages))
    else:
        messages.success(
            request,
            f"{count} materiais importados. Repetir esta confirmação não duplica os registros.",
        )
    return redirect("portability:preview", job_id=job_id)


@never_cache
@login_required
def table_export(request, extension):
    rows = export_material_rows(request.user)
    if extension == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Materiais"
        for row in rows:
            sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        return download(
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "materiais.xlsx",
        )
    output = StringIO()
    csv.writer(output, delimiter=";").writerows(rows)
    return download(
        "\ufeff" + output.getvalue(), "text/csv; charset=utf-8", "materiais.csv"
    )


@never_cache
@login_required
def archive(request):
    form = ArchiveForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        if not request.user.check_password(form.cleaned_data["password"]):
            form.add_error("password", "Senha incorreta.")
        else:
            try:
                if request.POST.get("action") == "export":
                    return download(
                        json.dumps(export_archive(request.user), ensure_ascii=False),
                        "application/json",
                        "davigurumi-completo.json",
                    )
                upload = form.cleaned_data["upload"]
                if not upload or upload.size > 50 * 1024 * 1024:
                    raise ValidationError("Envie um pacote de até 50 MB.")
                if not form.cleaned_data["acknowledge"]:
                    raise ValidationError(
                        "Confirme que leu as condições da restauração."
                    )
                package = json.loads(upload.read(50 * 1024 * 1024 + 1))
                count = import_archive(owner=request.user, package=package)
            except (ValueError, UnicodeError, ValidationError, RecursionError) as exc:
                form.add_error(
                    None,
                    (
                        exc
                        if isinstance(exc, ValidationError)
                        else "Pacote JSON inválido."
                    ),
                )
            else:
                messages.success(
                    request,
                    f"{count} registros restaurados. Os links antigos estão revogados.",
                )
                return redirect("dashboard")
    return render(request, "portability/archive.html", {"form": form})


@never_cache
@login_required
def template(request, extension):
    from .tables import COLUMNS

    rows = [
        list(COLUMNS),
        [
            "Fio de exemplo",
            "fio",
            "g",
            "0",
            "",
            "",
            "",
            "",
            "",
            "Algodão",
            "Médio",
            "2,5 mm",
            "0",
            "Exemplo; remova ou substitua antes de importar.",
        ],
    ]
    if extension == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "Materiais"
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        output = BytesIO()
        book.save(output)
        return download(
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "modelo-materiais.xlsx",
        )
    if extension != "csv":
        from django.http import Http404

        raise Http404
    output = StringIO()
    csv.writer(output, delimiter=";").writerows(rows)
    return download(
        "\ufeff" + output.getvalue(), "text/csv; charset=utf-8", "modelo-materiais.csv"
    )


@never_cache
@login_required
@require_POST
def discard(request, job_id):
    from django.contrib.auth import get_user_model
    from django.db import transaction

    with transaction.atomic():
        get_user_model().objects.select_for_update().get(pk=request.user.pk)
        job = get_object_or_404(
            ImportJob.objects.select_for_update(), pk=job_id, owner=request.user
        )
        if job.applied_at:
            messages.error(
                request,
                "Importação aplicada não é descartada por este controle. O histórico de estoque permanece preservado.",
            )
        else:
            job.delete()
            messages.success(
                request, "Prévia descartada; nenhum movimento de estoque foi alterado."
            )
    return redirect("portability:index")
