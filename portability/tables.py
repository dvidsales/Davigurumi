import csv
import re
import unicodedata
import uuid
import zipfile
from xml.etree.ElementTree import ParseError
from decimal import Decimal, InvalidOperation
from datetime import timedelta
from io import BytesIO, StringIO
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db import transaction
from django.shortcuts import get_object_or_404
from django.utils import timezone
from materials.forms import MaterialForm
from materials.models import Material
from materials.services import create_material
from .models import ImportJob

COLUMNS = (
    "nome",
    "tipo",
    "unidade",
    "quantidade",
    "custo_unitario",
    "marca",
    "cor",
    "codigo_cor",
    "tex",
)
KINDS = {
    "fio": "yarn",
    "tecido": "fabric",
    "enchimento": "filling",
    "acessorio": "accessory",
    "embalagem": "packaging",
    "outro": "other",
}


def normalize(value):
    value = (
        unicodedata.normalize("NFKD", str(value or ""))
        .encode("ascii", "ignore")
        .decode()
        .lower()
        .strip()
    )
    return re.sub(r"\s+", "_", value)


def decimal_text(value, locale):
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(str(value))
        if not number.is_finite():
            raise ValidationError("Número inválido.")
        return str(number)
    text = str(value or "0").strip()
    if locale == "pt-br":
        if not re.fullmatch(r"\d+(?:,\d+)?|\d{1,3}(?:\.\d{3})+(?:,\d+)?", text):
            raise ValidationError(
                "Número inválido para formato brasileiro. Use vírgula decimal ou escolha formato com ponto."
            )
        text = text.replace(".", "").replace(",", ".")
    elif not re.fullmatch(r"-?\d+(?:\.\d+)?", text):
        raise ValidationError("Número inválido para formato com ponto decimal.")
    try:
        return str(Decimal(text))
    except InvalidOperation as exc:
        raise ValidationError("Número inválido.") from exc


def read_table(upload, separator=";", sheet=""):
    if upload.size > 5 * 1024 * 1024:
        raise ValidationError("Arquivo excede 5 MB.")
    raw = upload.read(5 * 1024 * 1024 + 1)
    if upload.name.lower().endswith(".csv"):
        try:
            rows = []
            for row in csv.reader(
                StringIO(raw.decode("utf-8-sig")), delimiter=separator
            ):
                if len(rows) >= 2001 or len(row) > 40:
                    raise ValidationError(
                        "Limite de 2000 linhas/40 colunas por importação."
                    )
                rows.append(row)
        except (UnicodeError, csv.Error) as exc:
            raise ValidationError(
                "CSV inválido. Use UTF-8 e o separador escolhido."
            ) from exc
    elif upload.name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(BytesIO(raw)) as archive:
                if (
                    len(archive.infolist()) > 1000
                    or sum(info.file_size for info in archive.infolist())
                    > 30 * 1024 * 1024
                ):
                    raise ValidationError("Planilha excede o limite descompactado.")
            workbook = load_workbook(
                BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            if len(workbook.sheetnames) > 1 and not sheet:
                raise ValidationError(
                    "Escolha uma aba explicitamente: " + ", ".join(workbook.sheetnames)
                )
            selected = sheet or workbook.sheetnames[0]
            if selected not in workbook.sheetnames:
                raise ValidationError("Aba não encontrada.")
            if (workbook[selected].max_column or 0) > 40 or (
                workbook[selected].max_row or 0
            ) > 2001:
                workbook.close()
                raise ValidationError(
                    "Limite de 2000 linhas/40 colunas por importação."
                )
            rows = []
            for row in workbook[selected].iter_rows():
                if len(row) > 40:
                    raise ValidationError("Limite de 40 colunas por importação.")
                if any(cell.data_type == "f" for cell in row):
                    raise ValidationError(
                        "Planilha contém fórmulas. Exporte os valores antes de importar."
                    )
                rows.append([cell.value for cell in row])
                if len(rows) > 2001:
                    raise ValidationError("Limite de 2000 linhas por importação.")
            workbook.close()
        except ValidationError:
            raise
        except (ValueError, KeyError, zipfile.BadZipFile, OSError, ParseError) as exc:
            raise ValidationError("XLSX inválido ou corrompido.") from exc
    else:
        raise ValidationError("Use CSV ou XLSX, sem macros.")
    if not rows or len(rows) > 2001 or len(rows[0]) > 40:
        raise ValidationError("Tabela vazia ou acima de 2000 linhas/40 colunas.")
    headers = [normalize(value) for value in rows[0]]
    if len(headers) != len(set(headers)):
        raise ValidationError("Há nomes de colunas duplicados.")
    return headers, rows[1:]


def prepare_rows(*, owner, headers, rows, mapping, locale, job_id):
    missing = [
        mapping[field]
        for field in ("nome", "tipo", "unidade")
        if normalize(mapping[field]) not in headers
    ]
    if missing:
        raise ValidationError("Colunas obrigatórias ausentes: " + ", ".join(missing))
    accepted = []
    errors = []
    seen = set()
    for number, row in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in row):
            continue
        values = dict(zip(headers, row))
        get = lambda name: values.get(normalize(mapping.get(name, name)), "")
        try:
            kind = normalize(get("tipo"))
            kind = KINDS.get(kind, kind)
            payload = {
                "name": str(get("nome") or "").strip(),
                "kind": kind,
                "unit": str(get("unidade") or "").strip(),
                "initial_quantity": decimal_text(get("quantidade"), locale),
                "initial_unit_cost": (
                    decimal_text(get("custo_unitario"), locale)
                    if get("custo_unitario") not in (None, "")
                    else ""
                ),
                "brand": str(get("marca") or ""),
                "color": str(get("cor") or ""),
                "color_code": str(get("codigo_cor") or ""),
                "tex": (
                    decimal_text(get("tex"), locale)
                    if get("tex") not in (None, "")
                    else ""
                ),
                "request_key": str(uuid.uuid5(job_id, str(number))),
            }
            if any(
                value.lstrip().startswith(("=", "+", "@"))
                for value in (payload["name"], payload["brand"], payload["color"])
            ):
                raise ValidationError(
                    "Texto com aparência de fórmula não é aceito nesta importação."
                )
            form = MaterialForm(payload)
            if not form.is_valid():
                raise ValidationError(
                    [
                        str(message)
                        for messages in form.errors.values()
                        for message in messages
                    ]
                )
            identity = (payload["name"].casefold(), kind, payload["unit"])
            if (
                identity in seen
                or Material.objects.filter(
                    owner=owner,
                    name__iexact=payload["name"],
                    kind=kind,
                    unit=payload["unit"],
                ).exists()
            ):
                raise ValidationError(
                    "Material duplicado. Esta importação cria novos registros, sem atualizar cadastros silenciosamente."
                )
            seen.add(identity)
            accepted.append({"line": number, "payload": payload})
        except ValidationError as exc:
            errors.append({"line": number, "messages": exc.messages})
    return accepted, errors


def preview_import(*, owner, upload, mapping, locale, separator=";", sheet=""):
    headers, rows = read_table(upload, separator, sheet)
    job_id = uuid.uuid4()
    accepted, errors = prepare_rows(
        owner=owner,
        headers=headers,
        rows=rows,
        mapping=mapping,
        locale=locale,
        job_id=job_id,
    )
    return ImportJob.objects.create(
        id=job_id, owner=owner, headers=headers, rows=accepted, errors=errors
    )


@transaction.atomic
def confirm_import(*, owner, job_id):
    get_user_model().objects.select_for_update().get(pk=owner.pk)
    job = get_object_or_404(
        ImportJob.objects.select_for_update(), pk=job_id, owner=owner
    )
    if job.applied_at:
        return len(job.rows)
    if job.errors or not job.rows:
        raise ValidationError(
            "Resolva os erros antes de importar; nenhum registro foi alterado."
        )
    if timezone.now() - job.created_at > timedelta(days=2):
        raise ValidationError("Prévia expirada. Envie o arquivo novamente.")
    for entry in job.rows:
        payload = entry["payload"]
        if Material.objects.filter(
            owner=owner,
            name__iexact=payload["name"],
            kind=payload["kind"],
            unit=payload["unit"],
        ).exists():
            raise ValidationError(
                "Os cadastros mudaram após a prévia. Envie novamente para revisar duplicidades."
            )
        form = MaterialForm(payload)
        if not form.is_valid():
            raise ValidationError("Linha não é mais válida. Gere uma nova prévia.")
        create_material(form=form, owner=owner)
    job.applied_at = timezone.now()
    job.save(update_fields=["applied_at"])
    return len(job.rows)


def safe_cell(value):
    text = str(value)
    return (
        "'" + text
        if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r"))
        else text
    )


def export_material_rows(owner):
    rows = [list(COLUMNS)]
    for material in Material.objects.filter(owner=owner).prefetch_related(
        "layers", "movements"
    ):
        from materials.stock import reference_cost

        cost = reference_cost(material)
        kind = dict((value, key) for key, value in KINDS.items()).get(
            material.kind, material.kind
        )
        rows.append(
            [
                safe_cell(material.name),
                kind,
                material.unit,
                str(material.physical_stock),
                str(cost) if cost is not None else "",
                safe_cell(material.brand),
                safe_cell(material.color),
                safe_cell(material.color_code),
                str(material.tex) if material.tex is not None else "",
            ]
        )
    return rows
